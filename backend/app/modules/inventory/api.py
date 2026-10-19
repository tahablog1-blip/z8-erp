# modules/inventory/api.py — طبقة الـ HTTP: مسارات + صلاحيات فقط
# نفس عقد النظام القديم: GET /qty مفتوح عن قصد — نقطة البيع بتستعلم عن رصيد
# الصنف قبل البيع حتى لموظف بلا صلاحية عرض المخزون الكامل.
from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel

from ...core.deps import CurrentUser, get_current_user, require_permission
from . import service
from .schemas import (
    LevelOut, MoveOut, MoveResult, QtyOut, ReceiveBody, ShelfLocationIn,
    TransferBody, VoucherDetailOut, VoucherEditBody, VoucherOut,
)

router = APIRouter()


# ══════════════════ قراءة ══════════════════

@router.get("/levels", response_model=list[LevelOut])
async def list_levels(user: CurrentUser = Depends(require_permission("products.view_stock"))):
    """أرصدة كل الأصناف في المستودع وكل الفروع"""
    return await service.list_levels(user.company_id)


@router.get("/qty", response_model=QtyOut)
async def get_qty(locType: str = Query(...), locId: str = Query(...),
                  productId: str = Query(...),
                  user: CurrentUser = Depends(get_current_user)):
    """رصيد صنف واحد في موقع واحد — استعلام سريع لنقطة البيع"""
    return {"qty": await service.get_qty(user.company_id, locType, locId, productId)}


@router.get("/low-stock-alerts")
async def low_stock_alerts(user: CurrentUser = Depends(require_permission("products.view_stock"))):
    """تنبيهات النواقص بتفصيل كل موقع (مستودع/فرع) — لا الرصيد الكلي المجمّع"""
    return await service.low_stock_alerts(user.company_id)


@router.get("/suppliers", response_model=list[str])
async def list_suppliers(user: CurrentUser = Depends(get_current_user)):
    """أسماء الموردين المستخدمة سابقاً — اقتراحات لحقل المورد"""
    return await service.list_suppliers(user.company_id)


@router.get("/moves", response_model=list[MoveOut])
async def list_moves(limit: int = Query(default=100, ge=1, le=300),
                     user: CurrentUser = Depends(require_permission("products.view_stock"))):
    """سجل الحركات الفردية، أحدث أولاً"""
    return await service.list_moves(user.company_id, limit)


@router.get("/vouchers", response_model=list[VoucherOut])
async def list_vouchers(user: CurrentUser = Depends(require_permission("products.view_stock"))):
    """السندات مجمّعة حسب رقم السند"""
    return await service.list_vouchers(user.company_id)


@router.get("/vouchers/{move_no}", response_model=VoucherDetailOut)
async def get_voucher(move_no: str,
                      user: CurrentUser = Depends(require_permission("products.view_stock"))):
    return await service.get_voucher(user.company_id, move_no)


# ══════════════════ كتابة ══════════════════

@router.post("/receive", response_model=MoveResult, status_code=201)
async def receive(body: ReceiveBody,
                  user: CurrentUser = Depends(
                      require_permission("purchasing.invoices", "products.view_stock"))):
    """استلام من مورد خارجي إلى المستودع (سند واحد بعدة أصناف)"""
    try:
        return await service.receive(
            user.company_id, user.user_id, user.full_name,
            [i.model_dump() for i in body.items], body.supplierName, body.note,
        )
    except RuntimeError as e:
        # شجرة الحسابات ناقصة أو القيد غير متوازن — رسالة واضحة بدل 500 غامض
        raise HTTPException(400, str(e))


@router.post("/transfer", response_model=MoveResult, status_code=201)
async def transfer(body: TransferBody,
                   user: CurrentUser = Depends(require_permission("products.view_stock"))):
    """تحويل داخلي بين أي موقعين (مستودع↔فرع أو فرع↔فرع)"""
    return await service.transfer(
        user.company_id, user.user_id, user.full_name, body.model_dump()
    )


@router.put("/shelf-location")
async def set_shelf_location(body: ShelfLocationIn,
                             user: CurrentUser = Depends(require_permission("products.view_stock"))):
    """تحديد/تعديل موقع الرف الداخلي لصنف في موقع معيّن (مستودع أو فرع)"""
    return await service.set_shelf_location(
        user.company_id, body.locType, body.locId, body.productId, body.shelfLocation)


@router.put("/vouchers/{move_no}", response_model=MoveResult)
async def edit_voucher(move_no: str, body: VoucherEditBody,
                       user: CurrentUser = Depends(require_permission("products.view_stock"))):
    """تعديل سند: استبدال قائمة أصنافه بالكامل مع تصحيح المخزون تلقائياً"""
    return await service.edit_voucher(
        user.company_id, user.user_id, user.full_name, move_no,
        [i.model_dump() for i in body.items],
    )


# ══════════════════ طلبات الفروع (Requisitions) ══════════════════

class RequestItemIn(BaseModel):
    productId: str
    productLabel: str
    qty: int


class BranchRequestIn(BaseModel):
    requestingBranchId: str | None = None   # الفرع الطالب — إجباري لو المستخدم مش مربوط بفرع (زي الأدمن)
    sourceBranchId: str | None = None       # معرّف الفرع أو "main" للمستودع الرئيسي
    note: str | None = None
    items: list[RequestItemIn]


class DecideRequestIn(BaseModel):
    approve: bool
    sourceBranchId: str | None = None
    quantities: dict[str, int] = {}
    note: str | None = None


@router.post("/requests", status_code=201)
async def create_request(body: BranchRequestIn,
                         user: CurrentUser = Depends(require_permission("inventory.request"))):
    """مدير الفرع يطلب أصناف — بدون تنفيذ مباشر، لحد ما حد يوافق"""
    requesting_branch_id = body.requestingBranchId or user.branch_id
    if not requesting_branch_id:
        raise HTTPException(422, "اختر الفرع الطالب — حسابك غير مربوط بفرع افتراضي")
    return await service.create_request(
        user.company_id, requesting_branch_id, user.user_id, user.full_name, body.model_dump())


@router.get("/requests")
async def list_requests(branchId: str | None = Query(default=None),
                        status: str | None = Query(default=None),
                        user: CurrentUser = Depends(require_permission(
                            "inventory.request", "inventory.approve_requests"))):
    """لو معاه صلاحية الموافقة يشوف كل الفروع، غير كده يشوف طلبات فرعه بس"""
    effective_branch = branchId if user.has("inventory.approve_requests") else user.branch_id
    return await service.list_requests(user.company_id, effective_branch, status)


@router.post("/requests/{request_id}/decide")
async def decide_request(request_id: str, body: DecideRequestIn,
                         user: CurrentUser = Depends(require_permission("inventory.approve_requests"))):
    """موافقة (كلية/جزئية) تنفّذ التحويل تلقائياً، أو رفض بسبب"""
    return await service.decide_request(
        user.company_id, user.user_id, user.full_name, request_id, body.model_dump())


@router.get("/reorder-suggestions")
async def reorder(user: CurrentUser = Depends(require_permission("products.view_stock", "purchasing.invoices"))):
    """اقتراحات إعادة الطلب — الأصناف عند/تحت الحد الأدنى"""
    return await service.reorder_suggestions(user.company_id)


# ══════════════════ الجرد الفعلي ══════════════════

class StocktakeCreateBody(BaseModel):
    locType: str
    locId: str
    note: str | None = None


class StocktakeCountBody(BaseModel):
    countedQty: int
    note: str | None = None


def _check_stocktake_location(user: CurrentUser, loc_type: str, loc_id: str):
    """الأدمن لكل المواقع؛ الموظف يجب أن يكون مربوطاً بفرع محدد."""
    if user.role == "admin":
        return
    if not user.branch_scope:
        raise HTTPException(403, "حسابك غير مربوط بفرع؛ راجع مدير النظام")
    if loc_type != "branch" or str(loc_id) != str(user.branch_scope):
        raise HTTPException(403, "لا يمكنك الوصول إلى جرد هذا الموقع")


async def _check_stocktake_scope(user: CurrentUser, session_id: str) -> dict:
    """التحقق من أن جلسة الجرد ضمن نطاق المستخدم."""
    session = await service.get_stocktake(user.company_id, session_id)
    _check_stocktake_location(user, session["locType"], session["locId"])
    return session


@router.get("/stocktakes")
async def list_stocktakes_api(
    locType: str | None = Query(default=None),
    user: CurrentUser = Depends(require_permission("inventory.stocktake")),
):
    """قائمة جلسات الجرد."""
    if user.role != "admin" and not user.branch_scope:
        raise HTTPException(403, "حسابك غير مربوط بفرع؛ راجع مدير النظام")
    rows = await service.list_stocktakes(user.company_id, locType)
    if user.role != "admin":
        rows = [
            row for row in rows
            if row["locType"] == "branch"
            and str(row["locId"]) == str(user.branch_scope)
        ]
    return rows


@router.post("/stocktakes", status_code=201)
async def create_stocktake_api(
    body: StocktakeCreateBody,
    user: CurrentUser = Depends(require_permission("inventory.stocktake")),
):
    """إنشاء جلسة جرد وأخذ لقطة للأرصدة."""
    _check_stocktake_location(user, body.locType, body.locId)
    return await service.create_stocktake(
        user.company_id, user.user_id, user.full_name,
        body.locType, body.locId, body.note,
    )


@router.get("/stocktakes/{session_id}")
async def get_stocktake_api(
    session_id: str,
    user: CurrentUser = Depends(require_permission("inventory.stocktake")),
):
    """تفاصيل الجلسة والأصناف والكميات والفروق."""
    return await _check_stocktake_scope(user, session_id)


@router.put("/stocktakes/{session_id}/items/{product_id}")
async def count_stocktake_item_api(
    session_id: str,
    product_id: str,
    body: StocktakeCountBody,
    user: CurrentUser = Depends(require_permission("inventory.stocktake")),
):
    """تسجيل العدّ الفعلي دون تغيير رصيد المخزون."""
    await _check_stocktake_scope(user, session_id)
    return await service.count_stocktake_item(
        user.company_id, session_id, product_id,
        body.countedQty, body.note,
    )


@router.post("/stocktakes/{session_id}/review")
async def submit_stocktake_for_review_api(
    session_id: str,
    user: CurrentUser = Depends(require_permission("inventory.stocktake")),
):
    """إرسال جلسة الجرد للمراجعة."""
    await _check_stocktake_scope(user, session_id)
    return await service.submit_stocktake_for_review(user.company_id, session_id)


@router.post("/stocktakes/{session_id}/approve")
async def approve_stocktake_api(
    session_id: str,
    user: CurrentUser = Depends(require_permission("inventory.stocktake_approve")),
):
    """اعتماد الجرد وتسوية الأرصدة بعد التحقق من عدم تغيرها."""
    await _check_stocktake_scope(user, session_id)
    return await service.approve_stocktake(
        user.company_id, user.user_id, user.full_name, session_id,
    )


# ══════════ إلغاء الجلسة + رفع ملف الجرد ══════════
from fastapi import File, UploadFile  # noqa: E402

_BARCODE_HEADERS = {"barcode", "الباركود", "باركود"}
_MAX_UPLOAD = 10 * 1024 * 1024


def _parse_barcodes(filename: str, data: bytes) -> dict:
    """يقرأ العمود الأول فقط؛ كل صف = قطعة، والمكرر يُجمع."""
    import csv
    import io
    from collections import Counter

    name = (filename or "").lower()
    cells: list = []
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="ignore")
        cells = [r[0] if r else None for r in csv.reader(io.StringIO(text))]
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
            cells = [r[0] for r in ws.iter_rows(min_col=1, max_col=1, values_only=True)]
            wb.close()
        except Exception:
            raise HTTPException(400, "تعذّر قراءة الملف؛ تأكد أنه Excel صالح")
    else:
        raise HTTPException(400, "الصيغة المسموحة: xlsx أو csv")

    out: Counter = Counter()
    for v in cells:
        if v is None:
            continue
        if isinstance(v, float) and v.is_integer():
            v = int(v)  # يمنع 6.28E+12 و .0
        s = str(v).strip()
        if not s or s.lower() in _BARCODE_HEADERS:
            continue
        out[s] += 1
    return dict(out)


@router.post("/stocktakes/{session_id}/cancel")
async def cancel_stocktake_api(
    session_id: str,
    user: CurrentUser = Depends(require_permission("inventory.stocktake")),
):
    """إلغاء جلسة جرد غير معتمدة."""
    await _check_stocktake_scope(user, session_id)
    return await service.cancel_stocktake(user.company_id, session_id)


@router.post("/stocktakes/{session_id}/import")
async def import_stocktake_api(
    session_id: str,
    file: UploadFile = File(...),
    commit: bool = Query(default=False),
    zeroRest: bool = Query(default=False),
    user: CurrentUser = Depends(require_permission("inventory.stocktake")),
):
    """معاينة (commit=false) أو حفظ (commit=true) كميات الجرد من ملف باركودات."""
    await _check_stocktake_scope(user, session_id)
    data = await file.read()
    if len(data) > _MAX_UPLOAD:
        raise HTTPException(413, "الملف أكبر من 10 ميجا")
    scanned = _parse_barcodes(file.filename or "", data)
    return await service.import_stocktake_counts(
        user.company_id, session_id, scanned, commit, zeroRest,
    )
